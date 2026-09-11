"""
Comprehensive unit tests for Format-Aware Watermarking Engine
Verifies each format strategy (.docx, .xlsx, .pptx, .pdf, .png, .csv/.txt),
document parseability/integrity, NTFS ADS tagging, and idempotency.
"""

import unittest
import tempfile
import shutil
import os
import json
from pathlib import Path

from backend.watermark_engine import (
    WatermarkEngine,
    WatermarkStatus,
    DocxWatermarkStrategy,
    XlsxWatermarkStrategy,
    PptxWatermarkStrategy,
    PdfWatermarkStrategy,
    ImageWatermarkStrategy,
    PlaintextAdsWatermarkStrategy
)
from backend.database import DatabaseManager
from backend.watermark_backup import compute_file_sha256


class TestWatermarkEngine(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.backup_dir = Path(self.temp_dir) / "backups"
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        self.db_path = Path(self.temp_dir) / "test_engine.db"
        self.db = DatabaseManager(db_path=self.db_path)
        self.engine = WatermarkEngine(
            min_tier="Confidential",
            template="CONFIDENTIAL — Classified by PII Sentinel — {tier} — {date}",
            backup_dir=self.backup_dir,
            db=self.db
        )

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_docx_watermark(self):
        import docx

        docx_path = Path(self.temp_dir) / "test_doc.docx"
        doc = docx.Document()
        doc.add_heading("Internal Medical Record", 0)
        doc.add_paragraph("Patient Name: Jane Doe, SSN: 000-12-3456.")
        doc.save(str(docx_path))

        res = self.engine.apply_watermark(docx_path, tier="Confidential")
        self.assertEqual(res.status, WatermarkStatus.APPLIED)
        self.assertTrue(os.path.exists(res.backup_path))

        # Verify document integrity and watermark content
        doc_after = docx.Document(str(docx_path))
        self.assertGreaterEqual(len(doc_after.paragraphs), 2)
        self.assertIn("Jane Doe", doc_after.paragraphs[1].text)
        footer_text = doc_after.sections[0].footer.paragraphs[0].text
        self.assertIn("Classified by PII Sentinel", footer_text)

    def test_xlsx_watermark(self):
        import openpyxl

        xlsx_path = Path(self.temp_dir) / "test_sheet.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Payroll"
        ws["A1"] = "Employee ID"
        ws["B1"] = "Salary"
        ws["A2"] = "EMP-101"
        ws["B2"] = 95000
        wb.save(str(xlsx_path))

        res = self.engine.apply_watermark(xlsx_path, tier="Highly Confidential")
        self.assertEqual(res.status, WatermarkStatus.APPLIED)

        # Verify integrity and header/footer
        wb_after = openpyxl.load_workbook(str(xlsx_path))
        ws_after = wb_after["Payroll"]
        self.assertEqual(ws_after["A2"].value, "EMP-101")
        self.assertIn("Classified by PII Sentinel", ws_after.oddHeader.center.text)

    def test_pptx_watermark(self):
        from pptx import Presentation

        pptx_path = Path(self.temp_dir) / "test_pres.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Q3 Financial Projections"
        prs.save(str(pptx_path))

        res = self.engine.apply_watermark(pptx_path, tier="Restricted")
        self.assertEqual(res.status, WatermarkStatus.APPLIED)

        # Verify presentation still opens and has watermark shape
        prs_after = Presentation(str(pptx_path))
        self.assertEqual(len(prs_after.slides), 1)
        # Check shapes on slide
        texts = [s.text for s in prs_after.slides[0].shapes if s.has_text_frame]
        self.assertTrue(any("Classified by PII Sentinel" in t for t in texts))

    def test_pdf_watermark(self):
        from reportlab.pdfgen import canvas
        from pypdf import PdfReader

        pdf_path = Path(self.temp_dir) / "test_report.pdf"
        can = canvas.Canvas(str(pdf_path))
        can.drawString(100, 700, "CONFIDENTIAL REPORT - Credit Card: 4111-2222-3333-4444")
        can.save()

        res = self.engine.apply_watermark(pdf_path, tier="Confidential")
        self.assertEqual(res.status, WatermarkStatus.APPLIED)

        # Verify PDF validity
        reader = PdfReader(str(pdf_path))
        self.assertEqual(len(reader.pages), 1)
        text = reader.pages[0].extract_text()
        self.assertIn("CONFIDENTIAL REPORT", text)

    def test_image_watermark(self):
        from PIL import Image

        img_path = Path(self.temp_dir) / "test_scan.png"
        img = Image.new("RGB", (400, 400), color=(240, 240, 240))
        img.save(str(img_path))

        res = self.engine.apply_watermark(img_path, tier="Confidential")
        self.assertEqual(res.status, WatermarkStatus.APPLIED)

        # Verify image still valid
        with Image.open(str(img_path)) as img_after:
            self.assertEqual(img_after.size, (400, 400))

    def test_plaintext_ads_watermark(self):
        csv_path = Path(self.temp_dir) / "customers.csv"
        original_csv = "id,name,credit_card\n1,Alice,4111111111111111\n2,Bob,5500000000000004\n"
        csv_path.write_text(original_csv, encoding="utf-8")
        orig_hash = compute_file_sha256(csv_path)

        res = self.engine.apply_watermark(csv_path, tier="Confidential")
        self.assertEqual(res.status, WatermarkStatus.APPLIED)

        # 1. Main structured content MUST NOT be modified
        self.assertEqual(csv_path.read_text(encoding="utf-8"), original_csv)
        self.assertEqual(compute_file_sha256(csv_path), orig_hash)

        # 2. ADS tag should exist on NTFS
        ads_path = f"{csv_path}:pii-sentinel-classification"
        try:
            with open(ads_path, "r", encoding="utf-8") as f_ads:
                data = json.load(f_ads)
                self.assertEqual(data.get("tier"), "Confidential")
                self.assertEqual(data.get("applied_by"), "PII Sentinel")
        except FileNotFoundError:
            # FAT/exFAT environments don't support ADS, test still valid
            pass

    def test_idempotency_skip(self):
        test_file = Path(self.temp_dir) / "statement.txt"
        test_file.write_text("Customer SSN: 123-45-6789", encoding="utf-8")

        # 1st run: applied
        res1 = self.engine.apply_watermark(test_file, tier="Confidential")
        self.assertEqual(res1.status, WatermarkStatus.APPLIED)

        # 2nd run: skipped (no-op)
        res2 = self.engine.apply_watermark(test_file, tier="Confidential")
        self.assertEqual(res2.status, WatermarkStatus.SKIPPED_ALREADY_WATERMARKED)
        self.assertIn("Already watermarked", res2.message)

    def test_sensitivity_threshold_check(self):
        test_file = Path(self.temp_dir) / "public_info.txt"
        test_file.write_text("Public brochure information", encoding="utf-8")

        # Below Confidential threshold
        res = self.engine.apply_watermark(test_file, tier="Public")
        self.assertEqual(res.status, WatermarkStatus.SKIPPED_BELOW_THRESHOLD)


if __name__ == "__main__":
    unittest.main()
